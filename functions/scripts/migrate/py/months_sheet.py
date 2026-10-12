import json, collections, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
R='migration/raw/'
P=json.load(open(R+'replay-plan.json')); plan=P['plan']
S=json.load(open(R+'replay-state.json'))
students={s['id']:s for s in S['students']}
by_code={s['code']:s['id'] for s in S['students'] if s['code']}
sid=lambda k: k[4:] if k.startswith('doc:') else by_code.get(k[5:])
nm=lambda i: f"{students[i]['firstName'] or ''} {students[i]['lastName'] or ''}".strip()
now=P['createdAt']
F=lambda **k: Font(**{'name':'Arial','size':10,**k})
HEAD=PatternFill('solid',fgColor='1F2937'); INP=PatternFill('solid',fgColor='FFFF00'); GREY=PatternFill('solid',fgColor='F3F4F6'); ORG=PatternFill('solid',fgColor='FED7AA'); GRN=PatternFill('solid',fgColor='DCFCE7'); BLU=PatternFill('solid',fgColor='DBEAFE')
thin=Side(style='thin',color='D1D5DB'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def sgt(utcmin): return dt.datetime.fromisoformat(utcmin)+dt.timedelta(hours=8)
rows=collections.defaultdict(list)
for l in plan['lessons']:
    s=sid(l['studentKey']); t=sgt(l['startUTC'])
    past=l['startUTC']+':00Z' < now
    if l['marked']: status,fill='Present (marked in Classcard)',GRN
    elif past: status,fill='Not marked in Classcard',ORG
    else: status,fill='Upcoming',BLU
    rows[t.strftime('%Y-%m')].append((t, nm(s), students[s]['code'] or '', 'Online' if l['online'] else 'In person', 'Piano (Allie)' if l['teacher']=='piano' else 'Daniel', status, 'Consumed 1 session' if l['marked'] and past else ('Booked, waiting for you to mark' if past else 'Booked'), fill))
for u in plan['unidentified']:
    t=sgt(u['startUTC'])
    rows[t.strftime('%Y-%m')].append((t,'(no student recorded in Classcard)','', 'Online' if u['online'] else 'In person','Daniel','Present (marked in Classcard)' if u['marked'] else 'Not marked in Classcard','Not in LessonKaki (no student): add by hand if it was a real lesson',GREY))
wb=Workbook(); ws0=wb.active; ws0.title='Read me'
notes=['Month-by-month review (personal data, stays on your computer)','',
 'One tab per month. Every Classcard lesson on the calendar is listed, in time order.',
 'Green = Classcard had it marked present: LessonKaki has it as attended, 1 session used.',
 'Orange = past lesson Classcard never marked: it is in the LessonKaki calendar as booked, not used. Mark it present (or cancel it) after checking WhatsApp.',
 'Blue = still to come. Grey = a slot with no student on it in Classcard (time off, or a cancelled session whose name was removed): not in LessonKaki.',
 'Bright yellow column: type Y if it really happened and is not yet in LessonKaki, N if not. It is just your working list.',
 '',
 'In LessonKaki: Lessons > "Past, not marked" (or the calendar) to mark a lesson present; the credit comes off the right package automatically.']
for i,n in enumerate(notes,1):
    c=ws0.cell(i,1,n); c.font=F(bold=(i==1),size=13 if i==1 else 10); c.alignment=Alignment(wrap_text=True)
ws0.column_dimensions['A'].width=120
H=['Date','Time (SGT)','Student','Code','Where','Teacher','Classcard says','In LessonKaki','Happened? (Y/N)','Notes']
for m in sorted(rows):
    ws=wb.create_sheet(m)
    for c,h in enumerate(H,1):
        x=ws.cell(1,c,h); x.font=F(bold=True,color='FFFFFF'); x.fill=HEAD; x.border=BOX
    for r,(t,n,code,where,teacher,status,inlk,fill) in enumerate(sorted(rows[m],key=lambda x:(x[0],x[1])),2):
        vals=[t.strftime('%a %d %b'),t.strftime('%H:%M'),n,code,where,teacher,status,inlk,None,None]
        for c,v in enumerate(vals,1):
            x=ws.cell(r,c,v); x.font=F(); x.border=BOX
            x.fill = INP if c in (9,10) else fill
    for c,w in enumerate([12,10,32,10,10,13,30,46,14,30],1): ws.column_dimensions[get_column_letter(c)].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=f"A1:J{len(rows[m])+1}"
wb.save(R+'replay-months.xlsx')
print({m:len(v) for m,v in sorted(rows.items())})
