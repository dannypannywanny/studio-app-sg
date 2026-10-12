import json, csv, collections, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

R = 'migration/raw/'
P = json.load(open(R + 'replay-plan.json'))
plan, keys, attended = P['plan'], P['studentKeys'], P['attended']
S = json.load(open(R + 'replay-state.json'))
NOW = dt.datetime.fromisoformat(P['createdAt'].replace('Z', '+00:00'))
sgt = lambda iso: (dt.datetime.fromisoformat(iso.replace('Z', '+00:00')) + dt.timedelta(hours=8)).strftime('%Y-%m-%d') if iso else ''
sgt_dt = lambda iso: (dt.datetime.fromisoformat(iso.replace('Z', '+00:00')) + dt.timedelta(hours=8)).strftime('%a %d %b %Y %H:%M')

students = {s['id']: s for s in S['students']}
by_code = {s['code']: s['id'] for s in S['students'] if s['code']}
def sid(key):
    return key[4:] if key.startswith('doc:') else by_code.get(key[5:])
name = lambda i: f"{students[i]['firstName'] or ''} {students[i]['lastName'] or ''}".strip() if i in students else '?'
code = lambda i: (students[i]['code'] or '') if i in students else ''

# invoice amounts/dates from the activity file
act = {}
for l in open(R + 'classcard-activity.tsv'):
    if l.startswith('#') or l.startswith('ccId'): continue
    p = l.rstrip('\n').split('\t')
    act[p[3]] = {'date': p[4], 'amount': p[5], 'status': p[6]}

lots_state = {l['id']: l for l in S['lots']}
lessons_by_inv = collections.defaultdict(list)
for l in plan['lessons']:
    if l['invoiceId']: lessons_by_inv[l['invoiceId']].append(l)
unlinked = {sid(u['studentKey']): u['count'] for u in plan['unlinked']}

# ---- per student ----
marked_past = collections.Counter()
for l in plan['lessons']:
    if l['marked'] and l['startUTC'] < NOW.strftime('%Y-%m-%dT%H:%M'): marked_past[sid(l['studentKey'])] += 1
attended_by = collections.Counter()
for cc, n in attended.items():
    k = keys.get(cc)
    if k: attended_by[sid(k)] += n
lots_by_student = collections.defaultdict(list)
for l in plan['lots']: lots_by_student[sid(l['studentKey'])].append(l)
expired_by = collections.Counter(); remaining_by = collections.Counter()
for l in S['lots']:
    expired_by[l['studentId']] += l['expired']
past_unmarked = collections.Counter(); upcoming = collections.Counter(); last = {}; nxt = {}
nowiso = NOW.isoformat().replace('+00:00', 'Z')
for l in S['lessons']:
    if l['status'] == 'cancelled': continue
    s = l['studentId']
    if l['start'] < nowiso:
        if l['state'] != 'marked': past_unmarked[s] += 1
        if s not in last or l['start'] > last[s]: last[s] = l['start']
    else:
        upcoming[s] += 1
        if s not in nxt or l['start'] < nxt[s]: nxt[s] = l['start']
owed_by = collections.Counter(o['studentId'] for o in S['owed'])

# ---- styles ----
F = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k})
HEAD = PatternFill('solid', fgColor='1F2937'); RED = PatternFill('solid', fgColor='FECACA'); ORG = PatternFill('solid', fgColor='FED7AA')
YEL = PatternFill('solid', fgColor='FEF08A'); INP = PatternFill('solid', fgColor='FFFF00'); GREY = PatternFill('solid', fgColor='F3F4F6'); BLU = PatternFill('solid', fgColor='DBEAFE')
thin = Side(style='thin', color='D1D5DB'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(wb, title, headers, rows, widths, fills=None, inputs=(), note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = F(italic=True); ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(headers), 8)); ws.row_dimensions[1].height = 48
        r0 = 2
    for c, h in enumerate(headers, 1):
        x = ws.cell(r0, c, h); x.font = F(bold=True, color='FFFFFF'); x.fill = HEAD
        x.alignment = Alignment(wrap_text=True, vertical='center'); x.border = BOX
        if c - 1 in inputs: x.fill = PatternFill('solid', fgColor='A16207')
    for i, row in enumerate(rows):
        for c, v in enumerate(row, 1):
            x = ws.cell(r0 + 1 + i, c, v); x.font = F(); x.border = BOX; x.alignment = Alignment(vertical='top', wrap_text=isinstance(v, str) and len(v) > 40)
            if c - 1 in inputs: x.fill = INP
        if fills and fills[i]:
            for c in range(1, len(headers) + 1):
                if c - 1 not in inputs: ws.cell(r0 + 1 + i, c).fill = fills[i]
    for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 3)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return ws

wb = Workbook(); wb.remove(wb.active)

# ===== Balances =====
brows, bfills = [], []
order = []
for i in sorted(set(lots_by_student) | set(owed_by) | set(unlinked) | set(past_unmarked), key=lambda x: name(x).lower()):
    if i not in students: continue
    extra_marked = marked_past[i] - attended_by.get(i, 0)
    short = max(0, unlinked.get(i, 0) - S['adjust'].get(i, 0))
    flags, sev = [], 0
    if owed_by[i]: flags.append(f"{owed_by[i]} lesson(s) owed: marked with no package covering it"); sev = max(sev, 3)
    if extra_marked > 0: flags.append(f"Attendance mismatch: replay marked {extra_marked} more than Classcard says"); sev = max(sev, 2)
    if short: flags.append(f"{short} undated attended lesson(s) couldn't be deducted (no credit left)"); sev = max(sev, 2)
    if expired_by[i]: flags.append(f"{expired_by[i]} unused session(s) lost to expiry"); sev = max(sev, 1)
    if S['adjust'].get(i): flags.append(f"{S['adjust'][i]} undated attended lesson(s) deducted")
    order.append((-sev, name(i).lower(), i, flags))
for _, _, i, flags in sorted(order):
    sev = 3 if owed_by[i] else 2 if any('mismatch' in f or "couldn't" in f for f in flags) else 1 if expired_by[i] else 0
    bfills.append(RED if sev == 3 else ORG if sev == 2 else YEL if sev == 1 else None)
    brows.append([name(i), code(i), students[i]['status'], students[i]['available'], owed_by[i], expired_by[i], S['adjust'].get(i, 0),
                  sum(l['sessions'] for l in lots_by_student[i]), past_unmarked[i], upcoming[i], sgt(last.get(i)), sgt(nxt.get(i)),
                  '; '.join(flags), None, None])
ws = sheet(wb, 'Balances', ['Student', 'Code', 'Export status', 'Sessions left (replay)', 'Owed lessons', 'Lost to expiry', 'Undated attended, deducted',
                           'Bought (sessions, since May 2025 data)', 'Past, not marked', 'Upcoming', 'Last lesson', 'Next lesson', 'Why it is highlighted',
                           'YOUR sessions left', 'Notes'],
           brows, [26, 10, 10, 11, 9, 9, 12, 14, 9, 9, 12, 12, 70, 12, 30], bfills, inputs=(13, 14),
           note='One row per student with packages or lessons in the replay. Red = lessons owed; orange = attendance to check; yellow = sessions lost to expiry. '
                'Type what YOU think they have left in the yellow column where it differs from "Sessions left (replay)".')

# ===== Packages to confirm =====
prow, pfill = [], []
def lot_outcome(l):
    s = lots_state.get('cc-' + l['key'])
    if not s: return ('', 0, 0)
    return (f"used {s['used']}, expired unused {s['expired']}, left {s['remaining']}", s['expired'], s['remaining'])
for l in plan['lots']:
    s = sid(l['studentKey']); res, exp, rem = lot_outcome(l)
    inv = l['invoiceNumber']; a = act.get(inv, {})
    base = l['invoiceId']
    booked = lessons_by_inv.get(base, [])
    if l['sizeSource'].startswith('lessons booked'): pri, fill, why = 1, RED, "Size unknown: I used the lessons booked on it. Probably an old (Calendly-era / pre-price-rise) package."
    elif l['sizeSource'] == 'price rule (check)': pri, fill, why = 2, ORG, 'Dec-2025 S$150/S$200 promo price: I assumed 4 sessions'
    elif l['validityAssumed'] and (exp or rem): pri, fill, why = 3, YEL, f"Valid-until assumed ({'30' if l['sessions'] <= 4 else '90'} days); sessions were left over when it lapsed"
    else: continue
    prow.append([pri, name(s), code(s), inv or f"(id {base})", a.get('date') or sgt(l['validFrom']), a.get('amount'), l['sessions'], l['sizeSource'],
                 sgt(l['validUntil']) if l['validUntil'] else 'no expiry', len(booked), sum(1 for b in booked if b['marked']), res, why, None, None, None])
    pfill.append(fill)
idx = sorted(range(len(prow)), key=lambda i: (prow[i][0], prow[i][1].lower(), prow[i][4]))
prow = [prow[i] for i in idx]; pfill = [pfill[i] for i in idx]
sheet(wb, 'Packages to confirm', ['Priority', 'Student', 'Code', 'Invoice no.', 'Invoice date', 'Paid (S$)', 'Sessions I used', 'How sized', 'Valid until I used',
                                 'Lessons booked on it', 'of which marked', 'Result after replay', 'Why it is here', 'ACTUAL sessions bought', 'ACTUAL valid until (YYYY-MM-DD)', 'Notes'],
      prow, [8, 24, 10, 14, 12, 9, 10, 20, 14, 10, 10, 34, 55, 14, 16, 30], pfill, inputs=(13, 14, 15),
      note='Priority 1 (red) = I could not size it; 2 (orange) = the Dec-2025 promo price rule; 3 (yellow) = I assumed the expiry and sessions were left over. '
           'Fill the yellow columns where you know the real number; leave the rest and I will keep my assumption.')

# ===== Owed lessons =====
orow, ofill = [], []
plan_lots_state = collections.defaultdict(list)
for l in S['lots']: plan_lots_state[l['studentId']].append(l)
for o in sorted(S['owed'], key=lambda o: (name(o['studentId']).lower(), o['start'])):
    ls = plan_lots_state[o['studentId']]
    if not ls: why = 'No package in the data for this student'
    elif not any(x['validUntil'] >= o['start'] for x in ls):
        d = (dt.datetime.fromisoformat(o['start'].replace('Z', '+00:00')) - dt.datetime.fromisoformat(max(x['validUntil'] for x in ls).replace('Z', '+00:00'))).days
        why = f"Every package had ended ({d} days before this lesson)"
    elif not any(x['validFrom'] <= o['start'] for x in ls): why = 'No package had started yet'
    else: why = 'Packages valid, but all sessions already used up'
    orow.append([name(o['studentId']), code(o['studentId']), sgt_dt(o['start']), why, '; '.join(f"{x['name']} ({sgt(x['validFrom'])} to {sgt(x['validUntil']) if x['validUntil'] < '9000' else 'no expiry'}, {x['total']} sessions)" for x in sorted(ls, key=lambda x: x['validFrom'])), None])
    ofill.append(None)
sheet(wb, 'Owed lessons', ['Student', 'Code', 'Lesson (Singapore time)', 'Why no credit was available', "The student's packages in the replay", 'What to do (e.g. extend valid-until, add old package, ok to leave owed)'],
      orow, [24, 10, 24, 44, 90, 40], ofill, inputs=(5,),
      note='Lessons Classcard marked present where the replay had no package to charge. They show as "owed" in LessonKaki and settle automatically when a package is sold or credit is added.')

# ===== Attendance mismatch =====
mrow = []
link_note = {'class': 'found by class', 'created~': 'near created time', 'created+order': 'matched by order', 'start': 'matched by start time', 'created+start': ''}
for i in sorted(students, key=lambda x: name(x).lower()):
    extra = marked_past[i] - attended_by.get(i, 0)
    if extra > 0 and attended_by.get(i, 0) >= 0:
        ds = [f"{sgt(l['startUTC'])} ({link_note[l['link']]})" for l in plan['lessons'] if sid(l['studentKey']) == i and l['marked'] and l['link'] != 'created+start']
        mrow.append([name(i), code(i), attended_by.get(i, 0), marked_past[i], extra, '; '.join(sorted(ds)), None])
sheet(wb, 'Attendance to check', ['Student', 'Code', "Classcard says attended", 'Replay marked present', 'Replay has more by', 'Marked lessons that were linked less certainly', 'Which is right / notes'],
      mrow, [24, 10, 12, 12, 10, 90, 40], [ORG] * len(mrow), inputs=(6,),
      note="Students where the replay marks more lessons present than Classcard's own attendance report. A lesson may have been linked to the wrong student, or Classcard's report counts differently. Dates listed are the less certain links.")

# ===== Read me =====
ws = wb.create_sheet('Read me', 0)
lines = [
    ('LessonKaki: lessons replay reconciliation', True),
    ('Personal data: stays on your computer (migration/raw is not in git). Do not forward.', False),
    ('', False),
    ('What this is', True),
    ('Every student\'s balance in LessonKaki is rebuilt from Classcard history: packages bought, lessons booked, lessons Classcard marked present. Below are the places I am unsure.', False),
    ('', False),
    ('Colours', True),
    ('Red = something is clearly wrong or unknown (lessons owed, package size unknown).', False),
    ('Orange = please check (attendance mismatch, promo price rule).', False),
    ('Yellow = I made an assumption that only matters if it is wrong.', False),
    ('Bright yellow cells = for you to type into. Leave blank to keep my assumption.', False),
    ('', False),
    ('Tabs', True),
    ('Balances: everyone, with what the replay says they have left. Type your own number in "YOUR sessions left" where it differs.', False),
    ('Packages to confirm: invoices I could not size, or assumed. Fill in actual sessions and valid-until.', False),
    ('Owed lessons: marked lessons that had no package to charge, and why.', False),
    ('Attendance to check: students whose replayed attendance is higher than Classcard\'s own report.', False),
    ('', False),
    ('Not in this file (by your decision)', True),
    ('Events with no student (time off, cancelled sessions Classcard emptied): left as they are. For students whose attendance Classcard counts but which sit on those events, the count is deducted (column "Undated attended, deducted").', False),
    ('', False),
    ('Counts', True),
    (f"{len(S['students'])} students in LessonKaki, {len(plan['lots'])} packages, {len(plan['lessons'])} lessons replayed. Expired packages are not carried over.", False),
]
for r, (t, b) in enumerate(lines, 1):
    c = ws.cell(r, 1, t); c.font = F(bold=b, size=13 if r == 1 else 10); c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 120
for r, col in ((8, RED), (9, ORG), (10, YEL), (11, INP)): ws.cell(r, 1).fill = col

out = R + 'replay-reconcile.xlsx'
wb.save(out)
print('saved', out, {s.title: s.max_row for s in wb.worksheets})
