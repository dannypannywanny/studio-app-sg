# Turns the owner's annotated month sheets (replay-months-*.xlsx) into an actions file for
# apply-review.ts. Run from the repo root with a Python that has openpyxl:
#   python functions/scripts/migrate/py/review_actions.py ~/Downloads/sheet.xlsx 2026-04 2026-09 migration/raw/actions.json
# Answers in column "Happened? (Y/N)": Y / "Y - ..." = present; "CANCELLATION BY ME" / "CANCELLED BY ME" =
# teacher cancel; "CANCELLATION BY STUDENT" / "CANCELLED BY STUDENT" = student cancel. Rows with no student
# (no code) can't be applied and are only counted. Personal data stays in migration/raw (gitignored).
import collections, datetime as dt, json, sys
from openpyxl import load_workbook

xlsx, first, last, out = sys.argv[1:5]
R = 'migration/raw/'
plan = json.load(open(R + 'replay-plan.json'))['plan']
S = json.load(open(R + 'replay-state.json'))
by_code = {s['code']: s['id'] for s in S['students'] if s['code']}
sid = lambda k: k[4:] if k.startswith('doc:') else by_code.get(k[5:])
codes = {s['id']: s['code'] for s in S['students']}
idx = {}
for l in plan['lessons']:
    t = dt.datetime.fromisoformat(l['startUTC']) + dt.timedelta(hours=8)
    idx[(t.strftime('%Y-%m'), t.strftime('%a %d %b'), t.strftime('%H:%M'), codes[sid(l['studentKey'])])] = l['id']
acts, skipped, unknown = [], collections.Counter(), []
for ws in load_workbook(xlsx).worksheets:
    if not (first <= ws.title <= last):
        continue
    for r in list(ws.iter_rows(values_only=True))[1:]:
        d, t, name, code, where, teacher, cc, lk, ans = r[:9]
        if not ans:
            continue
        a = str(ans).strip().upper()
        if not code:
            skipped[a] += 1
            continue
        lid = idx.get((ws.title, d, t, code))
        if not lid:
            unknown.append((ws.title, d, t, code))
        elif a.startswith('Y'):
            acts.append({'lessonId': lid, 'action': 'present'})
        elif 'BY ME' in a:
            acts.append({'lessonId': lid, 'action': 'cancel', 'reason': 'Cancelled by teacher'})
        elif 'BY STUDENT' in a:
            acts.append({'lessonId': lid, 'action': 'studentCancel'})
        else:
            unknown.append(('?', a))
json.dump(acts, open(out, 'w'))
print(collections.Counter(a['action'] for a in acts), 'no-student skipped', dict(skipped), 'unmatched', len(unknown))
